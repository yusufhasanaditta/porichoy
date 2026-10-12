"""Render an improved CV as an ATS-friendly, single-column Word document."""

import io
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ACCENT = RGBColor(0x00, 0x6A, 0x4E)


def _heading(doc: Any, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = ACCENT


def _dates(start: str | None, end: str | None) -> str:
    return " – ".join(x for x in (start, end) if x)


def build_docx(cv: dict[str, Any]) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    for section in doc.sections:
        section.left_margin = section.right_margin = Pt(54)
        section.top_margin = section.bottom_margin = Pt(48)

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = name.add_run(cv["name"])
    r.bold = True
    r.font.size = Pt(20)
    doc.add_paragraph(cv["headline"]).runs[0].font.size = Pt(12)
    contact = cv.get("contact") or {}
    line = " · ".join(v for v in (contact.get(k) for k in ("email", "phone", "location", "linkedin", "website")) if v)
    if line:
        doc.add_paragraph(line)

    if cv.get("summary"):
        _heading(doc, "Summary")
        doc.add_paragraph(cv["summary"])

    if cv.get("experience"):
        _heading(doc, "Experience")
        for job in cv["experience"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.add_run(f"{job['title']} — {job['company']}").bold = True
            meta = " · ".join(x for x in (job.get("location"), _dates(job.get("start"), job.get("end"))) if x)
            if meta:
                doc.add_paragraph(meta).runs[0].italic = True
            for b in job["bullets"]:
                doc.add_paragraph(b, style="List Bullet")

    if cv.get("projects"):
        _heading(doc, "Projects")
        for proj in cv["projects"]:
            doc.add_paragraph().add_run(proj["name"]).bold = True
            for b in proj["bullets"]:
                doc.add_paragraph(b, style="List Bullet")

    if cv.get("education"):
        _heading(doc, "Education")
        for ed in cv["education"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.add_run(f"{ed['degree']} — {ed['school']}").bold = True
            meta = " · ".join(x for x in (_dates(ed.get("start"), ed.get("end")), ed.get("details")) if x)
            if meta:
                doc.add_paragraph(meta)

    if cv.get("skills"):
        _heading(doc, "Skills")
        for group in cv["skills"]:
            p = doc.add_paragraph()
            p.add_run(f"{group['category']}: ").bold = True
            p.add_run(", ".join(group["items"]))

    for title, key in (("Certifications", "certifications"), ("Languages", "languages")):
        if cv.get(key):
            _heading(doc, title)
            doc.add_paragraph(", ".join(cv[key]))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
